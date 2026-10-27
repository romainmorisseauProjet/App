import os
import sqlite3
import base64
import shutil
import subprocess
from datetime import datetime
import actualites
import musique
import carburants
import iss
import projets
import calcul

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

# Dossier où vivent TOUTES les données (base SQLite, musique, conversions...).
# Sur Render, pointe DATA_DIR vers le disque persistant (ex. /var/data) :
# sans disque, tout ce qui est dans data/ disparaît à chaque redéploiement.
DATA_DIR = os.environ.get("DATA_DIR") or os.path.join(BASE_DIR, "data")
DB_PATH = os.path.join(DATA_DIR, "assistant.db")
UI_DIR = os.path.join(BASE_DIR, "ui")
AUDIO_PATH = os.path.join(DATA_DIR, "audio")
CONVERSIONS_DIR = os.path.join(DATA_DIR, "conversions")
UPLOADS_DIR = os.path.join(DATA_DIR, "uploads")
# MUSIQUE_DIR permet de pointer vers ton dossier de musique existant (en local).
MUSIQUE_DIR = os.environ.get("MUSIQUE_DIR") or os.path.join(DATA_DIR, "musique")


def init_db():
    """Crée le fichier de base de données et les tables nécessaires
    si elles n'existent pas encore. Les futurs modules (manga, canevas IA,
    etc.) ajouteront leurs propres tables ici au fur et à mesure."""
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    conn = sqlite3.connect(DB_PATH)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS meteo_favoris (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nom TEXT NOT NULL,
            region TEXT,
            pays TEXT,
            latitude REAL NOT NULL,
            longitude REAL NOT NULL,
            ordre INTEGER NOT NULL
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS domotique_appareils (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nom TEXT NOT NULL,
            type TEXT NOT NULL,
            piece TEXT,
            etat INTEGER NOT NULL DEFAULT 0,
            valeur INTEGER NOT NULL DEFAULT 0,
            ordre INTEGER NOT NULL
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS modules_ordre (
            cle TEXT PRIMARY KEY,
            ordre INTEGER NOT NULL
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS parametres (
            cle TEXT PRIMARY KEY,
            valeur TEXT
        )
    """)
    conn.commit()
    conn.close()
    projets.creer_tables(DB_PATH)
    calcul.creer_tables(DB_PATH)


class Api:
    """Toutes les méthodes publiques (sans _ au début) sont appelables par
    l'interface via window.pywebview.api.nom_methode(...). Dans la version
    web, bridge.js transforme cet appel en requête POST /api/nom_methode,
    que server.py redirige vers la méthode du même nom ici."""

    def ping(self):
        return "Connexion au serveur OK."

    def recevoir_audio(self, audio_base64):
        """Étape 2 : reçoit l'audio enregistré depuis l'interface et le
        sauvegarde sur le disque. Ne le transcrit pas encore (ça viendra
        à l'étape 3) — on vérifie juste que le son arrive bien jusqu'ici."""
        os.makedirs(AUDIO_PATH, exist_ok=True)
        if len(audio_base64) > 30_000_000:
            return "Enregistrement trop long."
        donnees = base64.b64decode(audio_base64)
        chemin = os.path.join(AUDIO_PATH, "dernier_enregistrement.webm")
        with open(chemin, "wb") as f:
            f.write(donnees)
        taille_ko = round(len(donnees) / 1024, 1)
        return f"Audio reçu : {taille_ko} Ko, enregistré dans data/audio/"

    # --- Favoris météo ---

    def get_favoris(self):
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT id, nom, region, pays, latitude, longitude FROM meteo_favoris ORDER BY ordre"
        ).fetchall()
        conn.close()
        return [dict(r) for r in rows]

    def add_favori(self, nom, region, pays, latitude, longitude):
        conn = sqlite3.connect(DB_PATH)
        existe = conn.execute(
            "SELECT 1 FROM meteo_favoris WHERE nom = ? AND latitude = ? AND longitude = ?",
            (nom, latitude, longitude),
        ).fetchone()
        if existe:
            conn.close()
            return {"ok": False, "message": "Déjà dans les favoris"}
        ordre = conn.execute("SELECT COALESCE(MAX(ordre), -1) + 1 FROM meteo_favoris").fetchone()[0]
        conn.execute(
            "INSERT INTO meteo_favoris (nom, region, pays, latitude, longitude, ordre) VALUES (?, ?, ?, ?, ?, ?)",
            (nom, region, pays, latitude, longitude, ordre),
        )
        conn.commit()
        conn.close()
        return {"ok": True}

    def remove_favori(self, favori_id):
        conn = sqlite3.connect(DB_PATH)
        conn.execute("DELETE FROM meteo_favoris WHERE id = ?", (favori_id,))
        conn.commit()
        conn.close()
        return {"ok": True}

    # --- Domotique ---
    # Pour l'instant les appareils sont gérés localement (mémorisés en base,
    # allumage/extinction et réglages simulés). Chaque type d'appareil est
    # préparé pour qu'on y branche plus tard un vrai protocole (Tuya, Philips
    # Hue, MQTT/Zigbee2MQTT, etc.) sans changer l'interface : il suffira
    # d'appeler la véritable API du fabricant dans toggle_appareil /
    # set_valeur_appareil au lieu de se contenter d'écrire en base.

    VALEURS_DEFAUT_PAR_TYPE = {"lampe": 100, "volet": 100, "thermostat": 20}

    def get_appareils(self):
        conn = sqlite3.connect(DB_PATH)
        conn.row_factory = sqlite3.Row
        rows = conn.execute(
            "SELECT id, nom, type, piece, etat, valeur FROM domotique_appareils ORDER BY ordre"
        ).fetchall()
        conn.close()
        return [dict(r) for r in rows]

    def add_appareil(self, nom, type_appareil, piece):
        nom = (nom or "").strip()
        piece = (piece or "").strip()
        if not nom:
            return {"ok": False, "message": "Le nom est obligatoire"}
        valeur = self.VALEURS_DEFAUT_PAR_TYPE.get(type_appareil, 0)
        conn = sqlite3.connect(DB_PATH)
        ordre = conn.execute(
            "SELECT COALESCE(MAX(ordre), -1) + 1 FROM domotique_appareils"
        ).fetchone()[0]
        cur = conn.execute(
            "INSERT INTO domotique_appareils (nom, type, piece, etat, valeur, ordre) "
            "VALUES (?, ?, ?, 0, ?, ?)",
            (nom, type_appareil, piece, valeur, ordre),
        )
        conn.commit()
        nouvel_id = cur.lastrowid
        conn.close()
        return {
            "ok": True,
            "appareil": {
                "id": nouvel_id, "nom": nom, "type": type_appareil,
                "piece": piece, "etat": 0, "valeur": valeur,
            },
        }

    def remove_appareil(self, appareil_id):
        conn = sqlite3.connect(DB_PATH)
        conn.execute("DELETE FROM domotique_appareils WHERE id = ?", (appareil_id,))
        conn.commit()
        conn.close()
        return {"ok": True}

    def toggle_appareil(self, appareil_id):
        conn = sqlite3.connect(DB_PATH)
        row = conn.execute(
            "SELECT etat FROM domotique_appareils WHERE id = ?", (appareil_id,)
        ).fetchone()
        if row is None:
            conn.close()
            return {"ok": False}
        nouvel_etat = 0 if row[0] else 1
        conn.execute(
            "UPDATE domotique_appareils SET etat = ? WHERE id = ?",
            (nouvel_etat, appareil_id),
        )
        conn.commit()
        conn.close()
        return {"ok": True, "etat": nouvel_etat}

    def set_valeur_appareil(self, appareil_id, valeur):
        """Utilisé par les curseurs (luminosité, température, ouverture du
        volet…). Régler une valeur rallume l'appareil, comme dans la plupart
        des applis domotique."""
        conn = sqlite3.connect(DB_PATH)
        conn.execute(
            "UPDATE domotique_appareils SET valeur = ?, etat = 1 WHERE id = ?",
            (valeur, appareil_id),
        )
        conn.commit()
        conn.close()
        return {"ok": True}

    # --- Ordre des modules (glisser-déposer dans l'onglet Modules) ---

    def get_ordre_modules(self):
        conn = sqlite3.connect(DB_PATH)
        rows = conn.execute("SELECT cle FROM modules_ordre ORDER BY ordre").fetchall()
        conn.close()
        return [r[0] for r in rows]

    def set_ordre_modules(self, cles):
        conn = sqlite3.connect(DB_PATH)
        conn.execute("DELETE FROM modules_ordre")
        conn.executemany(
            "INSERT INTO modules_ordre (cle, ordre) VALUES (?, ?)",
            [(cle, i) for i, cle in enumerate(cles)],
        )
        conn.commit()
        conn.close()
        return {"ok": True}

    # --- Conversion de fichiers ---
    # Images <-> PDF (Pillow / PyMuPDF), documents bureautiques -> PDF
    # (LibreOffice, marche pour Word/Excel/PowerPoint/texte), PDF -> Word
    # (pdf2docx), PDF -> PowerPoint (pages rendues en images, via PyMuPDF +
    # python-pptx) et fusion de PDF (pypdf). Chaque conversion écrit ses
    # résultats dans un sous-dossier horodaté de data/conversions/, jamais
    # sur les fichiers d'origine.

    # Version web : les fichiers sont envoyés par le navigateur (POST /upload,
    # voir server.py) et rangés dans data/uploads/<id>/<nom>. L'interface ne
    # manipule que ces chemins relatifs "<id>/<nom>" ; le résultat est proposé
    # en téléchargement (/download/...).

    def _chemin_upload(self, relatif):
        """Transforme un chemin relatif envoyé par le navigateur en chemin
        absolu, en refusant tout ce qui sort du dossier d'uploads."""
        racine = os.path.realpath(UPLOADS_DIR)
        complet = os.path.realpath(os.path.join(racine, str(relatif)))
        if os.path.commonpath([racine, complet]) != racine or not os.path.isfile(complet):
            raise ValueError("Fichier non autorisé ou introuvable.")
        return complet

    def _nouveau_dossier_sortie(self, prefixe, dossier_parent=None):
        # Le paramètre dossier_parent est ignoré : le navigateur ne doit jamais
        # pouvoir choisir un chemin d'écriture sur le serveur.
        os.makedirs(CONVERSIONS_DIR, exist_ok=True)
        horodatage = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
        dossier = os.path.join(CONVERSIONS_DIR, f"{horodatage}_{prefixe}")
        os.makedirs(dossier, exist_ok=True)
        return dossier

    def images_vers_pdf(self, chemins, dossier_sortie=None):
        if not chemins:
            return {"ok": False, "message": "Choisis au moins une image."}
        try:
            chemins = [self._chemin_upload(c) for c in chemins]
            from PIL import Image
            images = [Image.open(c).convert("RGB") for c in chemins]
            dossier = self._nouveau_dossier_sortie("images-vers-pdf", dossier_sortie)
            sortie = os.path.join(dossier, "images.pdf")
            images[0].save(sortie, save_all=True, append_images=images[1:])
            return {"ok": True, "dossier": dossier, "message": f"{len(images)} image(s) rassemblée(s) dans un PDF."}
        except Exception as e:
            return {"ok": False, "message": f"Échec de la conversion : {e}"}

    def pdf_vers_images(self, chemin_pdf, format_image, dossier_sortie=None):
        if not chemin_pdf:
            return {"ok": False, "message": "Choisis un fichier PDF."}
        try:
            chemin_pdf = self._chemin_upload(chemin_pdf)
            import pymupdf
            dossier = self._nouveau_dossier_sortie("pdf-vers-images", dossier_sortie)
            doc = pymupdf.open(chemin_pdf)
            for i, page in enumerate(doc):
                pix = page.get_pixmap(dpi=200)
                pix.save(os.path.join(dossier, f"page-{i + 1}.{format_image}"))
            nb_pages = doc.page_count
            doc.close()
            return {"ok": True, "dossier": dossier, "message": f"{nb_pages} page(s) exportée(s) en image."}
        except Exception as e:
            return {"ok": False, "message": f"Échec de la conversion : {e}"}

    def document_vers_pdf(self, chemin_document, dossier_sortie=None):
        """Convertit n'importe quel document bureautique en PDF via LibreOffice :
        Word (.docx/.doc/.odt/.rtf), Excel (.xlsx/.xls/.ods/.csv),
        PowerPoint (.pptx/.ppt/.odp) et texte brut (.txt)."""
        if not chemin_document:
            return {"ok": False, "message": "Choisis un document."}
        soffice = shutil.which("soffice") or shutil.which("libreoffice")
        if not soffice:
            return {
                "ok": False,
                "message": "LibreOffice est requis pour cette conversion et n'est pas installé sur le serveur. "
                           "Il faut un déploiement Docker avec LibreOffice (voir le Dockerfile du projet).",
            }
        try:
            chemin_document = self._chemin_upload(chemin_document)
            dossier = self._nouveau_dossier_sortie("document-vers-pdf", dossier_sortie)
            subprocess.run(
                [soffice, "--headless", "--norestore", "--convert-to", "pdf", "--outdir", dossier, chemin_document],
                check=True, timeout=120,
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            )
            return {"ok": True, "dossier": dossier, "message": "Document converti en PDF."}
        except Exception as e:
            return {"ok": False, "message": f"Échec de la conversion : {e}"}

    def pdf_vers_docx(self, chemin_pdf, dossier_sortie=None):
        if not chemin_pdf:
            return {"ok": False, "message": "Choisis un fichier PDF."}
        try:
            chemin_pdf = self._chemin_upload(chemin_pdf)
            from pdf2docx import Converter
            dossier = self._nouveau_dossier_sortie("pdf-vers-word", dossier_sortie)
            nom = os.path.splitext(os.path.basename(chemin_pdf))[0]
            sortie = os.path.join(dossier, f"{nom}.docx")
            convertisseur = Converter(chemin_pdf)
            convertisseur.convert(sortie)
            convertisseur.close()
            return {"ok": True, "dossier": dossier, "message": "PDF converti en document Word."}
        except Exception as e:
            return {"ok": False, "message": f"Échec de la conversion : {e}"}

    def pdf_vers_pptx(self, chemin_pdf, dossier_sortie=None):
        """Convertit un PDF en présentation PowerPoint : chaque page devient
        une diapositive pleine page (image). Le texte n'est pas modifiable
        dans PowerPoint ensuite — pour ça, utilise plutôt PDF → Word."""
        if not chemin_pdf:
            return {"ok": False, "message": "Choisis un fichier PDF."}
        try:
            chemin_pdf = self._chemin_upload(chemin_pdf)
            import pymupdf
            from pptx import Presentation
            from pptx.util import Emu

            dossier = self._nouveau_dossier_sortie("pdf-vers-powerpoint", dossier_sortie)
            doc = pymupdf.open(chemin_pdf)
            nb_pages = doc.page_count
            if nb_pages == 0:
                doc.close()
                return {"ok": False, "message": "Ce PDF ne contient aucune page."}

            dpi = 150
            presentation = Presentation()
            diapo_vierge = presentation.slide_layouts[6]

            for i, page in enumerate(doc):
                pix = page.get_pixmap(dpi=dpi)
                chemin_image = os.path.join(dossier, f"_page-{i + 1}.png")
                pix.save(chemin_image)
                if i == 0:
                    presentation.slide_width = Emu(int(pix.width / dpi * 914400))
                    presentation.slide_height = Emu(int(pix.height / dpi * 914400))
                diapositive = presentation.slides.add_slide(diapo_vierge)
                diapositive.shapes.add_picture(
                    chemin_image, 0, 0,
                    width=presentation.slide_width, height=presentation.slide_height,
                )
                os.remove(chemin_image)
            doc.close()

            nom = os.path.splitext(os.path.basename(chemin_pdf))[0]
            sortie = os.path.join(dossier, f"{nom}.pptx")
            presentation.save(sortie)
            return {"ok": True, "dossier": dossier, "message": f"{nb_pages} page(s) transformée(s) en diapositives."}
        except Exception as e:
            return {"ok": False, "message": f"Échec de la conversion : {e}"}

    def fusionner_pdf(self, chemins, dossier_sortie=None):
        if not chemins or len(chemins) < 2:
            return {"ok": False, "message": "Choisis au moins deux fichiers PDF."}
        try:
            chemins = [self._chemin_upload(c) for c in chemins]
            from pypdf import PdfReader, PdfWriter
            writer = PdfWriter()
            for chemin in chemins:
                reader = PdfReader(chemin)
                for page in reader.pages:
                    writer.add_page(page)
            dossier = self._nouveau_dossier_sortie("pdf-fusionnes", dossier_sortie)
            sortie = os.path.join(dossier, "fusion.pdf")
            with open(sortie, "wb") as f:
                writer.write(f)
            return {"ok": True, "dossier": dossier, "message": f"{len(chemins)} PDF fusionnés en un seul fichier."}
        except Exception as e:
            return {"ok": False, "message": f"Échec de la fusion : {e}"}

    # --- Actualités ---

    def get_actualites(self, categorie):
        return actualites.recuperer_articles(categorie)

    # --- Musique ---
    # La bibliothèque vit sur le serveur (data/musique/). La lecture se fait
    # dans le navigateur (balise <audio>), qui récupère les fichiers via
    # /media/musique/<fichier> (voir server.py).

    def _lire_parametre(self, cle, defaut=None):
        conn = sqlite3.connect(DB_PATH)
        row = conn.execute("SELECT valeur FROM parametres WHERE cle = ?", (cle,)).fetchone()
        conn.close()
        return row[0] if row else defaut

    def _ecrire_parametre(self, cle, valeur):
        conn = sqlite3.connect(DB_PATH)
        conn.execute(
            "INSERT INTO parametres (cle, valeur) VALUES (?, ?) "
            "ON CONFLICT(cle) DO UPDATE SET valeur = excluded.valeur",
            (cle, valeur),
        )
        conn.commit()
        conn.close()

    def get_bibliotheque(self):
        os.makedirs(MUSIQUE_DIR, exist_ok=True)
        return musique.scanner_dossier(MUSIQUE_DIR)

    def telecharger_musique(self, url, format_cible):
        os.makedirs(MUSIQUE_DIR, exist_ok=True)
        return musique.telecharger(url, format_cible, MUSIQUE_DIR)

    # --- Carburants ---

    def get_stations_carburant(self, latitude, longitude, carburant, rayon_km):
        return carburants.rechercher_stations(latitude, longitude, carburant, rayon_km)

    # --- Station spatiale (ISS) ---

    def get_iss_etat(self):
        """Position + télémétrie + pays survolé, en un seul appel."""
        pos = iss.position_actuelle()
        if pos.get("ok"):
            pos["pays"] = iss.pays_survole(pos["latitude"], pos["longitude"])
        return pos

    def get_iss_trace(self):
        return iss.trace_orbite()

    def get_iss_equipage(self):
        return iss.equipage()

    def get_iss_passages(self, latitude, longitude):
        return iss.prochains_passages(latitude, longitude)

    def get_iss_lieu(self):
        """Lieu d'observation mémorisé (JSON {nom, lat, lon}), ou None."""
        import json
        brut = self._lire_parametre("iss_lieu")
        return json.loads(brut) if brut else None

    def set_iss_lieu(self, nom, latitude, longitude):
        import json
        self._ecrire_parametre("iss_lieu", json.dumps({"nom": nom, "lat": latitude, "lon": longitude}))
        return {"ok": True}


    # --- Projets (bloc-notes + avancement) ---
    # Toute la logique est dans projets.py ; ces méthodes ne font que faire le pont.

    def projets_lister(self):
        return projets.lister(DB_PATH)

    def projets_creer(self, nom):
        return projets.creer(DB_PATH, nom)

    def projets_renommer(self, projet_id, nom):
        return projets.renommer(DB_PATH, projet_id, nom)

    def projets_supprimer(self, projet_id):
        return projets.supprimer(DB_PATH, projet_id)

    def projets_definir_progression(self, projet_id, valeur):
        return projets.definir_progression_manuelle(DB_PATH, projet_id, valeur)

    def projets_details(self, projet_id):
        return projets.details(DB_PATH, projet_id)

    def projets_page_lire(self, page_id):
        return projets.page_lire(DB_PATH, page_id)

    def projets_page_creer(self, projet_id, titre="Nouvelle page"):
        return projets.page_creer(DB_PATH, projet_id, titre)

    def projets_page_sauver(self, page_id, titre, contenu):
        return projets.page_sauver(DB_PATH, page_id, titre, contenu)

    def projets_page_supprimer(self, page_id):
        return projets.page_supprimer(DB_PATH, page_id)

    def projets_tache_ajouter(self, projet_id, texte):
        return projets.tache_ajouter(DB_PATH, projet_id, texte)

    def projets_tache_basculer(self, tache_id):
        return projets.tache_basculer(DB_PATH, tache_id)

    def projets_tache_modifier(self, tache_id, texte):
        return projets.tache_modifier(DB_PATH, tache_id, texte)

    def projets_tache_supprimer(self, tache_id):
        return projets.tache_supprimer(DB_PATH, tache_id)

    # --- Calcul (expressions, conversions, maths, formules, photo) ---
    # Toute la logique est dans calcul.py ; ces méthodes ne font que faire le pont.

    def calcul_catalogue(self):
        return calcul.catalogue()

    def calcul_demander(self, texte, image=None, detail=False):
        return calcul.demander(DB_PATH, texte, image, detail)

    def calcul_formule(self, formule_id, valeurs, enregistrer=False):
        return calcul.calculer_formule(DB_PATH, formule_id, valeurs, enregistrer)

    def calcul_historique_lister(self):
        return calcul.historique_lister(DB_PATH)

    def calcul_historique_supprimer(self, entree_id):
        return calcul.historique_supprimer(DB_PATH, entree_id)

    def calcul_historique_vider(self):
        return calcul.historique_vider(DB_PATH)

    def calcul_ia_statut(self):
        return calcul.ia_statut(DB_PATH)

    def calcul_ia_configurer(self, cle, modele):
        return calcul.ia_configurer(DB_PATH, cle, modele)

    def calcul_ia_retirer_cle(self):
        return calcul.ia_retirer_cle(DB_PATH)


if __name__ == "__main__":
    # Lancement en local (VS Code, bouton ▶ ou « python main.py ») :
    # ouvre la fenêtre de bureau comme avant. Sur Render, c'est server.py
    # qui est lancé par gunicorn et ce bloc n'est jamais exécuté.
    import lancer
    lancer.demarrer()
